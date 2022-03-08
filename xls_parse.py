#!/usr/bin/env python
import argparse
import sys
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token


def dbgp(msg):
    return
    print(msg)


WORKBOOK = None
WORKBOOK_DATA = None
LINES = []
PROCESSED_CELLS = set()
WORKSHEET_STACK = []

FAPP_XML_OBJ = "fapp_xml"
SPECIAL_FAPP_XML_CALL = {
    "C4": f"{FAPP_XML_OBJ}.get_counter_timer_freq()",
    "G4": f"{FAPP_XML_OBJ}.get_measured_time()",
    "G5": f"{FAPP_XML_OBJ}.get_node_name()",
    "G10": f"{FAPP_XML_OBJ}.get_process_no()",
    "G11": f"{FAPP_XML_OBJ}.get_cmg_no()",
    "G12": f"{FAPP_XML_OBJ}.get_measured_region()",
    "C10": f"{FAPP_XML_OBJ}.get_vector_length()",
}

# These cells are read by formulas and are "hardcoded" in the excel
# sheet (i.e. they are not read from anywhere else, e.g. from xml).
DATA_VALUES = {
    "C6",
    "C7",
    "C8",
    "C9",
    "C14",
    "C15",
    "C16",
    "C17",
    "Y4",
}

HEADER_ROW = 29
TOP_ROW = 30
BOTOM_ROW = 41
LEFT_COLUMN = 29
RIGHT_COLUMN = 334

INFIX_OP_MAP = {
    "=": "==",
    "^": "**",
}


def is_event_cell(cell_id):
    cell = cell_id_to_obj(cell_id)
    col = cell.col_idx
    row = cell.row
    result = (
        LEFT_COLUMN <= col
        and col <= RIGHT_COLUMN
        and TOP_ROW <= row
        and row <= BOTOM_ROW
    )
    return result


def full_cell_id(cell_id):
    dbgp(f"> full_cell_id({cell_id})")
    prefix = WORKSHEET_STACK[-1] if WORKSHEET_STACK else "report"
    if cell_id in WORKBOOK.defined_names:
        cell_id = WORKBOOK.defined_names[cell_id].value
    elif "!" not in cell_id:
        cell_id = f"{prefix}!{cell_id}"
    dbgp(f"< full_cell_id -> {cell_id}")
    return cell_id


def cell_obj_to_id(cell):
    return cell.parent.title + "!" + cell.coordinate


def cell_id_to_varname(cell_id):
    dbgp(f"> cell_id_to_varname({cell_id})")
    if ":" in cell_id:
        cells = cell_id_to_obj(cell_id)
        cell_ids = [cell_obj_to_id(cell[0]) for cell in cells]
        cell_vars = [cell_id_to_varname(cell_id) for cell_id in cell_ids]
        result = f"{', '.join(cell_vars)}"
    else:
        cell_id = full_cell_id(cell_id)
        result = cell_id.replace("!", "_").replace("$", "")
    dbgp(f"< cell_id_to_varname -> {result}")
    return result


def cell_id_to_obj(cell_id):
    dbgp(f"> cell_id_to_obj({cell_id})")
    cell_id = full_cell_id(cell_id)
    ws, cell = cell_id.split("!")
    cell = WORKBOOK[ws][cell]
    if isinstance(cell, MergedCell):
        print("WARNING: MergedCell!")
    dbgp(f"< cell_id_to_obj -> {cell}")
    return cell


def unknown_type_exception(token):
    msg = f"ERROR: Unknown type {token.type}"
    raise Exception(msg)


def unknown_subtype_exception(token):
    msg = f"ERROR: Unknown subtype {token.subtype} "
    msg += f"(of token type {token.type})"
    raise Exception(msg)


def unknown_func_exception(token):
    msg = f"ERROR: Unknown FUNC {token.value}"
    raise Exception(msg)


def assert_sep_comma(token):
    assert token.type == Token.SEP and token.value == ","


def assert_func_close(token):
    assert token.type == Token.FUNC and token.subtype == Token.CLOSE


def python_cmd_to_read_xml(cell_id):
    dbgp(f"> python_cmd_to_read_xml({cell_id})")
    result = None
    cell_id = cell_id.replace("$", "")
    if ":" in cell_id:
        cells = cell_id_to_obj(cell_id)
        cell_ids = [cell_obj_to_id(cell[0]) for cell in cells]
        results = [python_cmd_to_read_xml(cell_id) for cell_id in cell_ids]
        if any(results):
            result = f"{', '.join(results)}"
    else:
        cell_id = full_cell_id(cell_id)
        ws, cell = cell_id.split("!")
        if ws == "label":
            value = WORKBOOK[ws][cell].value
            result = f"'{value}'"
        elif ws == "data":
            if cell in SPECIAL_FAPP_XML_CALL:
                result = SPECIAL_FAPP_XML_CALL[cell]
            elif cell in DATA_VALUES:
                result = str(WORKBOOK["data"][cell].value)
            elif is_event_cell(cell_id):
                cell_obj = WORKBOOK[ws][cell]
                col = cell_obj.col_idx - 1
                event_name = WORKBOOK[ws][HEADER_ROW][col].value
                thread_id = cell_obj.row - HEADER_ROW - 1
                result = f"{FAPP_XML_OBJ}.get_event('{event_name}', {thread_id})"
    dbgp(f"< python_cmd_to_read_xml -> {result}")
    return result


def parse_operand(token):
    dbgp(f"> parse_operand({token})")
    if token.subtype == Token.RANGE:
        cell_id = full_cell_id(token.value)
        raw = python_cmd_to_read_xml(cell_id)
        if raw:
            result = raw
        else:
            cell_to_inst(cell_id)
            result = cell_id_to_varname(cell_id)
    elif token.subtype == Token.TEXT:
        result = token.value
    elif token.subtype == Token.NUMBER:
        result = token.value
    else:
        unknown_subtype_exception(token)
    dbgp(f"< parse_operand -> {result}")
    return result


def parse_if(tokens, cur):
    cond, cur = parse_tokens(tokens, cur + 1)
    assert_sep_comma(tokens[cur])
    true_val, cur = parse_tokens(tokens, cur + 1)
    assert_sep_comma(tokens[cur])
    false_val, cur = parse_tokens(tokens, cur + 1)
    assert_func_close(tokens[cur])
    result = f"({true_val}) if ({cond}) else ({false_val})"
    return result, cur


def parse_or(tokens, cur):
    ops, cur = parse_tokens(tokens, cur + 1)
    while tokens[cur].type == Token.SEP and tokens[cur].value == ",":
        tmp, cur = parse_tokens(tokens, cur + 1)
        if tmp != "":
            ops += f", {tmp}"
    assert_func_close(tokens[cur])
    result = f"any([{ops}])"
    return result, cur


def parse_count(tokens, cur):
    cells, cur = parse_tokens(tokens, cur + 1)
    assert_func_close(tokens[cur])
    result = f"sum(1 for e in [{cells}] if e !='')"
    return result, cur


def parse_sum(tokens, cur):
    terms, cur = parse_tokens(tokens, cur + 1)
    assert_func_close(tokens[cur])
    result = f"(xls_sum([{terms}]))"
    return result, cur


def parse_average(tokens, cur):
    terms, cur = parse_tokens(tokens, cur + 1)
    while tokens[cur].type == Token.SEP and tokens[cur].value == ",":
        tmp, cur = parse_tokens(tokens, cur + 1)
        if tmp != "":
            terms += f", {tmp}"
    assert_func_close(tokens[cur])
    result = f"(xls_sum([{terms}]) / len(xls_nonempty([{terms}])))"
    return result, cur


def parse_gl_lower(tokens, cur):
    args, cur = parse_tokens(tokens, cur + 1)
    while tokens[cur].type == Token.SEP and tokens[cur].value == ",":
        tmp, cur = parse_tokens(tokens, cur + 1)
        if tmp != "":
            args += f", {tmp}"
    assert_func_close(tokens[cur])
    result = f"vba_guard_limit_lower({args})"
    return result, cur


def parse_gl_upper(tokens, cur):
    args, cur = parse_tokens(tokens, cur + 1)
    while tokens[cur].type == Token.SEP and tokens[cur].value == ",":
        tmp, cur = parse_tokens(tokens, cur + 1)
        if tmp != "":
            args += f", {tmp}"
    assert_func_close(tokens[cur])
    result = f"vba_guard_limit_upper({args})"
    return result, cur


def parse_func(tokens, cur):
    dbgp(f"> parse_func({tokens} {cur} = {tokens[cur]})")
    token = tokens[cur]
    open_func_dict = {
        "IF(": parse_if,
        "OR(": parse_or,
        "COUNT(": parse_count,
        "SUM(": parse_sum,
        "AVERAGE(": parse_average,
        "GuardLimitLower(": parse_gl_lower,
        "GuardLimitUpper(": parse_gl_upper,
    }
    if token.subtype == Token.OPEN:
        parse_fn = open_func_dict[token.value]
        result, cur = parse_fn(tokens, cur)
    elif token.subtype == Token.CLOSE:
        result = None
    else:
        unknown_subtype_exception(token)
    dbgp(f"< parse_func -> {result, cur}")
    return result, cur


def parse_infix_op(token):
    dbgp(f"> parse_infix_op({token})")
    op_name = token.value
    if token.value in INFIX_OP_MAP:
        op_name = INFIX_OP_MAP[token.value]
    dbgp(f"< parse_infix_op -> {op_name}")
    return op_name


def parse_tokens(tokens, cur):
    dbgp(f"> parse_tokens({tokens}, {cur} = {tokens[cur]})")
    result = ""
    while cur < len(tokens):
        token = tokens[cur]
        if token.type == Token.OPERAND:
            result += parse_operand(token)
        elif token.type == Token.FUNC:
            tmp, cur = parse_func(tokens, cur)
            if tmp:
                result += tmp
            else:
                break
        elif token.type == Token.OP_IN:
            result += parse_infix_op(token)
        elif token.type == Token.PAREN:
            result += token.value
        elif token.type == Token.WSPACE:
            pass
        elif token.type == Token.SEP:
            break
        else:
            unknown_type_exception(token)
        cur += 1
    dbgp(f"< parse_tokens -> {result, cur}")
    return result, cur


def cell_to_inst(cell_id):
    dbgp(f"> cell_to_inst({cell_id})")
    cell_id = full_cell_id(cell_id)
    cell = cell_id_to_obj(cell_id)
    if isinstance(cell, tuple):
        for c in cell:
            cell_to_inst(cell_obj_to_id(c[0]))
    else:
        cell_var = cell_id_to_varname(cell_id)
        if cell_var in PROCESSED_CELLS:
            return
        WORKSHEET_STACK.append(cell.parent.title)
        tokens = Tokenizer(cell.value).items
        cell_val, _ = parse_tokens(tokens, 0)
        WORKSHEET_STACK.pop()
        line = f"{cell_var} = {cell_val}"
        PROCESSED_CELLS.add(cell_var)
        LINES.append(line)
        dbgp(f"< cell_to_inst -> APPEND: {line}")


def get_label(cell_id):
    cell = WORKBOOK_DATA["report"][cell_id]
    return cell.value


def create_program(output):
    result = []
    with open("fapp_loader.py") as loader:
        result += loader.readlines()

    with open("fapp_top.py.in") as top:
        result += top.readlines()

    result += [line + "\n" for line in LINES]

    with open("fapp_bottom.py.in") as top:
        result += top.readlines()

    result = "".join(result)
    with open(output, "w") as out:
        out.write(result)
    return result


def record_entry(prefix, key, value):
    LINES.append(
        f"add_path({[get_label(e) for e in prefix]}, '{key}', {value}, results)"
    )


def add_key_single_value_pair(prefix, key, value):
    cell_to_inst(value)
    record_entry(prefix, get_label(key), cell_id_to_varname(value))


def add_column_of_12_1(prefix, key, first, num_rows=12):
    first_cell = cell_id_to_obj(first)
    row = first_cell.row
    col = first_cell.col_idx - 1
    if isinstance(WORKBOOK["report"][row + 1][col], MergedCell):
        add_key_single_value_pair(prefix, key, first)
    else:
        cell_varnames = []
        for offset in range(num_rows):
            cell_id = cell_obj_to_id(WORKBOOK["report"][row + offset][col])
            cell_to_inst(cell_id)
            cell_varnames.append(cell_id_to_varname(cell_id))

        total_id = cell_obj_to_id(WORKBOOK["report"][row + num_rows][col])
        cell_to_inst(total_id)

        value = f"[e for e in [{', '.join(cell_varnames)}] if e != '']"
        label = get_label(key)
        record_entry(prefix, label, value)
        # total_var = cell_id_to_varname(total_id)
        # record_entry(prefix, f"{label}_total", total_var)
        pass


def _col2num(col, base=26):
    ords = list(map(lambda t: ord(t) - ord("A"), col))
    result = ords[0]
    if len(col) == 2:
        result = (ords[0] + 1) * base + ords[1]
    return result


def _num2col(num, base=26):
    q, r = divmod(num, base)
    result = chr(r + ord("A"))
    if q:
        result = chr(q - 1 + ord("A")) + result
    return result


def col_range(begin_col, end_col):
    assert len(begin_col) <= 2
    assert len(end_col) <= 2

    begin_num = _col2num(begin_col)
    end_num = _col2num(end_col)
    return [_num2col(num) for num in range(begin_num, end_num + 1)]


def add_table(prefix, begin_col, end_col, head_row, first_row):
    for col in col_range(begin_col, end_col):
        head_cell = col + str(head_row)
        first_cell = col + str(first_row)
        add_column_of_12_1(prefix, head_cell, first_cell)


def add_tables():
    # TOP
    add_key_single_value_pair([], "A3", "C3")
    add_key_single_value_pair([], "A4", "C4")
    add_key_single_value_pair([], "H3", "J3")
    # add_key_single_value_pair([], "H4", "J4")
    add_key_single_value_pair([], "H5", "J5")
    add_key_single_value_pair([], "O3", "Q3")
    add_key_single_value_pair([], "O4", "Q4")

    # STATISTICS
    add_table(["A8"], "C", "N", 8, 14)

    # CYCLE ACCOUNTING
    add_table(["P8", "R8"], "R", "S", 9, 14)
    add_table(["P8", "T8"], "T", "Y", 9, 14)
    add_table(["P8", "Z8"], "Z", "AA", 9, 14)
    add_table(["P8", "AB8"], "AB", "AC", 9, 14)
    add_table(["P8"], "AD", "AG", 8, 14)
    add_table(["P8", "AH8"], "AH", "AK", 9, 14)
    add_table(["P8"], "AL", "AL", 8, 14)

    # BUSY
    add_table(["A28"], "C", "P", 28, 34)

    # CACHE
    add_table(["A48"], "C", "P", 48, 55)

    # INSTRUCTIONS
    add_table(["A69", "C69", "C70", "C71"], "C", "I", 72, 77)
    add_table(["A69", "C69", "C70", "J71"], "J", "J", 72, 77)
    add_table(["A69", "C69", "K70", "K71"], "K", "O", 72, 77)
    add_table(["A69", "C69", "K70", "P71"], "P", "P", 72, 77)
    add_table(["A69", "Q69"], "Q", "S", 70, 77)
    add_table(["A69"], "T", "T", 69, 77)
    add_table(["A69", "U69"], "U", "W", 70, 77)
    add_table(["A69", "X69"], "X", "Y", 71, 77)
    add_table(["A69"], "Z", "AE", 69, 77)

    # POWER
    add_table(["AG69"], "AI", "AK", 69, 77)

    # PREFETCH
    add_table(["A92", "C92"], "C", "E", 93, 98)
    add_table(["A92", "F92"], "F", "H", 93, 98)
    add_table(["A92", "I92"], "I", "I", 93, 98)

    # FLOP
    add_table(["K92"], "M", "P", 92, 98)

    # EXTRA
    add_table(["R92", "T92"], "T", "V", 93, 98)
    add_table(["R92", "W92"], "W", "AB", 93, 98)
    add_table(["R92"], "AC", "AC", 92, 98)

    # # DATA TRANSFER
    # for col in colnames("C", "F"):
    #     add_key_single_value_pair(f"{col}113", f"{col}115")
    #     add_key_single_value_pair(f"{col}113", f"{col}116")


def main():
    global WORKBOOK
    global WORKBOOK_DATA

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "input_xls",
        type=str,
        help="Path to the input xls file",
    )
    args = parser.parse_args()
    filename = Path(args.input_xls).expanduser()

    WORKBOOK = openpyxl.load_workbook(filename)
    WORKBOOK_DATA = openpyxl.load_workbook(filename, data_only=True)

    add_tables()

    out_file = sys.argv[0].replace(".py", ".out.py")
    create_program(Path(out_file).expanduser())


main()
